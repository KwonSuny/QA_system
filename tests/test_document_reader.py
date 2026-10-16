"""验证真实文件解析、来源定位和上传失败时的写入边界。"""

import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import ArrayObject, DictionaryObject, NameObject, DecodedStreamObject

from document_reader import MAX_UPLOAD_BYTES, parse_document
from rag_pipeline import heading_contexts, normalize, source_documents


def docx_fixture():
    document = Document()
    document.add_heading("甲研究所", level=1)
    document.add_paragraph("甲研究所教师44人。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "数量"
    table.cell(1, 0).text = "实验室"
    table.cell(1, 1).text = "8个"
    document.add_paragraph("统计截至2024年。")
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def pdf_fixture(texts, encrypted=False):
    writer = PdfWriter()
    for text in texts:
        page = writer.add_blank_page(width=600, height=800)
        if not text:
            continue
        chars = list(dict.fromkeys(text))
        codes = {char: i + 1 for i, char in enumerate(chars)}
        mappings = "\n".join(f"<{codes[c]:04X}> <{c.encode('utf-16-be').hex()}>" for c in chars)
        cmap = DecodedStreamObject()
        cmap.set_data(("/CIDInit /ProcSet findresource begin 12 dict begin begincmap\n"
                       "/CIDSystemInfo << /Registry (Adobe) /Ordering (UCS) /Supplement 0 >> def\n"
                       "/CMapName /Test def /CMapType 2 def\n"
                       "1 begincodespacerange <0000> <FFFF> endcodespacerange\n"
                       f"{len(chars)} beginbfchar\n{mappings}\nendbfchar\n"
                       "endcmap CMapName currentdict /CMap defineresource pop end end").encode())
        descendant = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                       NameObject("/Subtype"): NameObject("/CIDFontType2"),
                                       NameObject("/BaseFont"): NameObject("/TestFont")})
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type0"),
                                 NameObject("/BaseFont"): NameObject("/TestFont"),
                                 NameObject("/Encoding"): NameObject("/Identity-H"),
                                 NameObject("/DescendantFonts"): ArrayObject([writer._add_object(descendant)]),
                                 NameObject("/ToUnicode"): writer._add_object(cmap)})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({
            NameObject("/F1"): writer._add_object(font)})})
        content = DecodedStreamObject()
        code = "".join(f"{codes[c]:04X}" for c in text)
        content.set_data(f"BT /F1 12 Tf 50 700 Td <{code}> Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(content)
    if encrypted:
        writer.encrypt("test-password")
    stream = io.BytesIO()
    writer.write(stream)
    return stream.getvalue()


class DocumentReaderTests(unittest.TestCase):
    def test_docx_order_table_heading_and_filename(self):
        parsed = parse_document("机构资料.DOCX", docx_fixture())
        documents = source_documents(parsed.text, "自定义库名")
        self.assertEqual(documents[0]["title"], "机构资料.DOCX")
        text = documents[0]["text"]
        self.assertLess(text.index("44人"), text.index("实验室"))
        self.assertLess(text.index("8个"), text.index("2024年"))
        self.assertIn("实验室 | 8个", text)
        sections = heading_contexts(text, [text])
        self.assertTrue(any("甲研究所" in s["heading"] for s in sections[0]))

    def test_pdf_chinese_pages_and_empty_page_warning(self):
        parsed = parse_document("机构.pdf", pdf_fixture(["甲研究所教师44人。", "乙研究所教师20人。", ""]))
        self.assertIn("甲研究所教师44人。", parsed.text)
        self.assertIn("## 第2页", parsed.text)
        self.assertIn("第3页", parsed.warnings[0])
        text = source_documents(parsed.text, "库")[0]["text"]
        sections = heading_contexts(text, ["甲研究所教师44人。", "乙研究所教师20人。"])
        self.assertTrue(any("第1页" in s["heading"] for s in sections[0]))
        self.assertTrue(any("第2页" in s["heading"] for s in sections[1]))

    def test_scanned_encrypted_corrupt_unsupported_empty_and_oversize(self):
        inputs = [("扫描.pdf", pdf_fixture([""]), "OCR"),
                  ("加密.pdf", pdf_fixture(["甲"], True), "加密"),
                  ("坏.pdf", b"not a pdf", "PDF"), ("坏.docx", b"not a zip", "DOCX"),
                  ("旧.doc", b"abc", "支持"), ("空.txt", b"", "为空"),
                  ("大.txt", b"a" * (MAX_UPLOAD_BYTES + 1), "20 MB")]
        for name, raw, error in inputs:
            with self.subTest(name=name), self.assertRaisesRegex(ValueError, error):
                parse_document(name, raw)

    def test_text_bom_gbk_and_existing_source_markers(self):
        text = "# 来源文档：已有来源\n中文资料"
        for encoding in ["utf-8-sig", "utf-16", "gbk"]:
            self.assertEqual(parse_document("资料.txt", text.encode(encoding)).text, text)

    def test_upload_route_and_failed_replacement(self):
        import server
        from fastapi import HTTPException, UploadFile
        manager = Mock()
        manager.create.return_value = {"id": "new", "chunk_count": 1}
        with patch.object(server, "get_manager", return_value=manager):
            result = server.upload_kb(UploadFile(filename="资料.pdf", file=io.BytesIO(pdf_fixture(["甲", ""]))),
                                      name="自定义", description="", kb_id="", user={"id": 7})
            self.assertIn("第2页", result["warnings"][0])
            self.assertEqual(manager.create.call_args.args[:2], (7, "自定义"))
            self.assertIn("# 来源文档：资料.pdf", manager.create.call_args.args[3])
            with self.assertRaises(HTTPException) as error:
                server.upload_kb(UploadFile(filename="资料.docx", file=io.BytesIO(b"bad")),
                                 name="", description="", kb_id="existing", user={"id": 7})
            self.assertEqual(error.exception.status_code, 400)
            manager.update.assert_not_called()

    def test_http_multipart_upload_and_unsupported_format(self):
        import server
        from fastapi.testclient import TestClient
        manager = Mock()
        manager.create.return_value = {"id": "new", "chunk_count": 1}
        previous = dict(server.app.dependency_overrides)
        server.app.dependency_overrides[server.get_current_user] = lambda: {"id": 9}
        try:
            client = TestClient(server.app)
            with patch.object(server, "get_manager", return_value=manager):
                response = client.post("/api/kbs/upload", data={"name": "机构资料"},
                                       files={"file": ("机构.docx", docx_fixture(),
                                                       "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.json()["filename"], "机构.docx")
                response = client.post("/api/kbs/upload", data={"kb_id": "existing"},
                                       files={"file": ("不支持.doc", b"abc", "application/msword")})
                self.assertEqual(response.status_code, 400)
                manager.update.assert_not_called()
        finally:
            server.app.dependency_overrides.clear()
            server.app.dependency_overrides.update(previous)

    def test_real_storage_retrieval_keeps_pdf_page_and_docx_table(self):
        from db import Database, User
        from kb_service import KnowledgeBaseManager
        root = Path(__file__).resolve().parents[1] / "kb_store"
        root.mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(prefix="document-reader-test-", dir=root) as directory:
            db = Database(Path(directory) / "test.db")
            manager = KnowledgeBaseManager(db, Path(directory) / "store")
            embed = Mock()
            embed.encode.side_effect = lambda values: [[0.2, 0.3, 1.0] for _ in ([values] if isinstance(values, str) else values)]
            manager._embedding_model = embed
            try:
                with db.session() as session:
                    user = User(username="document", salt="unused", password_hash="unused")
                    session.add(user)
                    session.flush()
                    uid = user.id
                for filename, raw, keyword in [("机构.pdf", pdf_fixture(["甲研究所教师44人。"]), "44人"),
                                               ("机构.docx", docx_fixture(), "8个")]:
                    kb = manager.create(uid, "自定义库名", text=parse_document(filename, raw).text)
                    catalog = manager.document_catalog(uid, kb["id"])
                    self.assertEqual(catalog[0]["title"], filename)
                    results = manager.query_evidence(uid, kb["id"], "D1", keyword, 4)
                    self.assertTrue(any(keyword in row["text"] for row in results))
                    if filename.endswith(".pdf"):
                        self.assertTrue(any("第1页" in s["heading"] for row in results for s in row["section_context"]))
                    saved = manager.get(uid, kb["id"])["text"]
                    self.assertIn(keyword, normalize(saved))
            finally:
                close = getattr(manager.db_client, "close", None)
                if close:
                    close()
                else:
                    manager.db_client._system.stop()
                db.engine.dispose()


if __name__ == "__main__":
    unittest.main()
