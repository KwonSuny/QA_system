"""将上传文件转换为可编辑、可索引的正文，不执行文件中的活动内容。"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import PurePosixPath

from rag_pipeline import merge_documents

MAX_UPLOAD_BYTES = 20 * 1024 * 1024
MAX_TEXT_CHARACTERS = 2_000_000
SUPPORTED_SUFFIXES = {".txt", ".md", ".text", ".docx", ".pdf"}


@dataclass
class ParsedDocument:
    text: str
    warnings: list[str] = field(default_factory=list)


def decode_text(raw: bytes) -> str:
    # BOM 必须先识别，UTF-16 的字节也可能碰巧符合 UTF-8。
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16")
    for encoding in ("utf-8-sig", "gbk", "utf-16"):
        try:
            return raw.decode(encoding)
        except UnicodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _docx_text(raw: bytes) -> ParsedDocument:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError("DOCX 解压后超过100 MB，请拆分后上传")
        document = Document(io.BytesIO(raw))
    except (zipfile.BadZipFile, KeyError, OSError) as exc:
        raise ValueError("DOCX 文件损坏或并非有效的 DOCX 文件") from exc
    parts = []

    def walk(container):
        for block in container.iter_inner_content():
            if isinstance(block, Paragraph):
                text = block.text.strip()
                if not text:
                    continue
                heading = re.fullmatch(r"Heading ([1-9])", block.style.name or "")
                parts.append(f"{'#' * int(heading[1])} {text}" if heading else text)
            elif isinstance(block, Table):
                # 按正文顺序保留行列，不猜测第一行一定是表头。
                parts.append("表格：")
                for row in block.rows:
                    cells, previous = [], None
                    for cell in row.cells:
                        if cell._tc is previous:
                            continue
                        previous = cell._tc
                        before = len(parts)
                        walk(cell)
                        cells.append(" / ".join(parts[before:]))
                        del parts[before:]
                    if any(cells):
                        parts.append(" | ".join(cells))
                parts.append("")

    walk(document)
    warnings = []
    if document.inline_shapes:
        warnings.append("DOCX 包含图片，本次只读取正文和表格中的文本，未识别图片文字。")
    return ParsedDocument("\n\n".join(parts), warnings)


def _pdf_text(raw: bytes) -> ParsedDocument:
    from pypdf import PdfReader

    if not raw.startswith(b"%PDF-"):
        raise ValueError("文件内容并非有效的 PDF 文件")
    reader = PdfReader(io.BytesIO(raw))
    if reader.is_encrypted:
        raise ValueError("暂不支持加密 PDF，请先解密后上传")
    if len(reader.pages) > 500:
        raise ValueError("PDF 超过500页，请拆分后上传")
    parts, empty = [], []
    for number, page in enumerate(reader.pages, 1):
        content = page.get_contents()
        if content is not None and len(content.get_data()) > 8 * 1024 * 1024:
            raise ValueError(f"PDF 第{number}页内容过大，请拆分或简化后上传")
        text = (page.extract_text(extraction_mode="layout", layout_mode_strip_rotated=False) or "").strip() if content is not None else ""
        if text:
            parts.append(f"## 第{number}页\n\n{text}")
        else:
            empty.append(number)
    if not parts:
        raise ValueError("PDF 没有可提取的文本，可能是扫描件；请先进行 OCR 文字识别后上传")
    warnings = []
    if empty:
        pages = "、".join(map(str, empty[:20])) + ("等" if len(empty) > 20 else "")
        warnings.append(f"PDF 第{pages}页未提取到文本，可能为空白或图片页；本次未进行 OCR。")
    return ParsedDocument("\n\n".join(parts), warnings)


def parse_document(filename: str, raw: bytes) -> ParsedDocument:
    filename = PurePosixPath(filename.replace("\\", "/")).name
    suffix = PurePosixPath(filename).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError("支持 .txt、.md、.text、.docx、.pdf 文件")
    if len(raw) > MAX_UPLOAD_BYTES:
        raise ValueError("文件超过20 MB，请拆分后上传")
    if not raw:
        raise ValueError("文件为空")
    try:
        if suffix == ".docx":
            result = _docx_text(raw)
        elif suffix == ".pdf":
            result = _pdf_text(raw)
        else:
            result = ParsedDocument(decode_text(raw))
    except ImportError as exc:
        raise ValueError("缺少文件解析依赖，请安装 requirements-web.txt 并重启服务") from exc
    except ValueError:
        raise
    except Exception as exc:
        # 不将解析库的内部路径和异常细节发送给浏览器。
        raise ValueError("文件解析失败，请确认文件完整且格式正确") from exc
    if not result.text.strip():
        raise ValueError("文件没有可读取的正文")
    if len(result.text) > MAX_TEXT_CHARACTERS:
        raise ValueError("提取正文超过200万字符，请拆分后上传")
    if suffix in {".docx", ".pdf"}:
        result.text = merge_documents([{"title": filename, "text": result.text}])
    return result
